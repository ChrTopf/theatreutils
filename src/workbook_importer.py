from typing import Any, Optional

from openpyxl import load_workbook

from block import Block

BLOCKS_SHEET_NAME = "Blocks"
REQUIRED_HEADERS = ("ID", "Part", "Scene", "Page", "Type", "Text")


class WorkbookImporter:
    """Reads the blocks back from a workbook written by WorkbookExporter.

    The workbook may have been edited by hand in the meantime, typically to fill
    in the `Audio Effect` and `Light` columns. Columns are therefore looked up by
    their header name, so their order does not matter and any additional columns
    are ignored.
    """

    def import_blocks(self, path: str) -> list[Block]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if BLOCKS_SHEET_NAME not in workbook.sheetnames:
                raise ValueError(f"The workbook '{path}' does not contain a '{BLOCKS_SHEET_NAME}' sheet.")
            rows = workbook[BLOCKS_SHEET_NAME].iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                return []
            column_indices = self._column_indices(header)
            return [self._to_block(row, column_indices) for row in rows if self._is_block_row(row, column_indices)]
        finally:
            workbook.close()

    @staticmethod
    def _column_indices(header: tuple) -> dict[str, int]:
        column_indices = {str(name).strip(): index for index, name in enumerate(header) if name is not None}
        missing = [name for name in REQUIRED_HEADERS if name not in column_indices]
        if missing:
            raise ValueError(f"The '{BLOCKS_SHEET_NAME}' sheet is missing the columns: {', '.join(missing)}")
        return column_indices

    @staticmethod
    def _is_block_row(row: tuple, column_indices: dict[str, int]) -> bool:
        return WorkbookImporter._cell(row, column_indices, "ID") is not None

    @staticmethod
    def _to_block(row: tuple, column_indices: dict[str, int]) -> Block:
        def cell(name: str) -> Any:
            return WorkbookImporter._cell(row, column_indices, name)

        return Block(
            id=int(cell("ID")),
            type=str(cell("Type") or ""),
            part=str(cell("Part") or ""),
            scene=int(cell("Scene") or 0),
            page=int(cell("Page") or 0),
            text=str(cell("Text") or ""),
            character=WorkbookImporter._optional_text(cell("Character")),
            direction=WorkbookImporter._optional_text(cell("Direction")),
            mute=WorkbookImporter._name_list(cell("Mute")),
            unmute=WorkbookImporter._name_list(cell("Unmute")),
            audio_effect=WorkbookImporter._optional_text(cell("Audio Effect")),
            light=WorkbookImporter._optional_text(cell("Light")),
        )

    @staticmethod
    def _cell(row: tuple, column_indices: dict[str, int], name: str) -> Any:
        index = column_indices.get(name)
        if index is None or index >= len(row):
            return None
        return row[index]

    @staticmethod
    def _optional_text(value: Any) -> Optional[str]:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _name_list(value: Any) -> list[str]:
        if value is None:
            return []
        return [name.strip() for name in str(value).split(",") if name.strip()]
