from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from model.actor import Actor
from model.block import Block

HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

# Cycled per scene so consecutive scenes always contrast with each other.
SCENE_FILL_COLORS = ("DDEBF7", "FCE4D6", "E2EFDA", "FFF2CC", "EAD1DC")

TYPE_FONTS = {
    "dialogue": Font(),
    "acting": Font(italic=True, color="595959"),
    "command": Font(bold=True, color="C00000"),
}

SCENE_BORDER = Border(top=Side(style="thin", color="808080"))
PART_BORDER = Border(top=Side(style="thick", color="000000"))

BLOCK_HEADERS = ("ID", "Part", "Scene", "Page", "Type", "Character", "Direction", "Text", "Mute", "Unmute")
CHARACTER_HEADERS = ("Mic", "Actor", "Character", "Actions")
TEXT_COLUMN_INDEX = BLOCK_HEADERS.index("Text") + 1
TYPE_COLUMN_INDEX = BLOCK_HEADERS.index("Type") + 1


class WorkbookExporter:
    """Exports the parsed dataset (actors + blocks) as a styled LibreOffice/Excel workbook.

    One worksheet lists the character/actor/mic assignments, the other lists every
    block in script order. Blocks are shaded per scene (and set off with a border
    on every act change) so a reviewer can quickly spot where one scene ends and
    the next begins while proof-reading the digitized script.
    """

    def export(self, actors: list[Actor], blocks: list[Block], output_path: str) -> None:
        workbook = Workbook()

        characters_sheet = workbook.active
        characters_sheet.title = "Characters"
        self._write_characters_sheet(characters_sheet, actors)

        blocks_sheet = workbook.create_sheet("Blocks")
        self._write_blocks_sheet(blocks_sheet, blocks)

        workbook.save(output_path)

    def _write_characters_sheet(self, sheet: Worksheet, actors: list[Actor]) -> None:
        sheet.append(CHARACTER_HEADERS)
        self._style_header_row(sheet, len(CHARACTER_HEADERS))

        for actor in actors:
            first_row = sheet.max_row + 1
            for character in actor.characters:
                sheet.append([actor.mic, actor.actor, character.name, character.actions])
            last_row = sheet.max_row
            if last_row > first_row:
                sheet.merge_cells(start_row=first_row, start_column=1, end_row=last_row, end_column=1)
                sheet.merge_cells(start_row=first_row, start_column=2, end_row=last_row, end_column=2)
            for column_index in (1, 2):
                cell = sheet.cell(row=first_row, column=column_index)
                cell.alignment = Alignment(vertical="center", horizontal="center")

        self._autosize_columns(sheet, CHARACTER_HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    def _write_blocks_sheet(self, sheet: Worksheet, blocks: list[Block]) -> None:
        sheet.append(BLOCK_HEADERS)
        self._style_header_row(sheet, len(BLOCK_HEADERS))

        previous_scene_key = None
        previous_part = None
        fill_color_index = -1

        for block in blocks:
            scene_key = (block.part, block.scene)
            scene_changed = scene_key != previous_scene_key
            part_changed = block.part != previous_part
            if scene_changed:
                fill_color_index = (fill_color_index + 1) % len(SCENE_FILL_COLORS)
                previous_scene_key = scene_key
            previous_part = block.part

            sheet.append([block.id, block.part, block.scene, block.page, block.type, block.character, block.direction, block.text,
                          ", ".join(block.mute), ", ".join(block.unmute)])
            row_number = sheet.max_row
            fill = PatternFill(start_color=SCENE_FILL_COLORS[fill_color_index], end_color=SCENE_FILL_COLORS[fill_color_index], fill_type="solid")
            border = PART_BORDER if part_changed else (SCENE_BORDER if scene_changed else None)

            for column_index in range(1, len(BLOCK_HEADERS) + 1):
                cell = sheet.cell(row=row_number, column=column_index)
                cell.fill = fill
                if border is not None:
                    cell.border = border

            sheet.cell(row=row_number, column=TYPE_COLUMN_INDEX).font = TYPE_FONTS.get(block.type, Font())
            sheet.cell(row=row_number, column=TEXT_COLUMN_INDEX).alignment = Alignment(wrap_text=True, vertical="top")

        self._autosize_columns(sheet, BLOCK_HEADERS, wide_column_index=TEXT_COLUMN_INDEX)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    @staticmethod
    def _style_header_row(sheet: Worksheet, column_count: int) -> None:
        for column_index in range(1, column_count + 1):
            cell = sheet.cell(row=1, column=column_index)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 20

    @staticmethod
    def _autosize_columns(sheet: Worksheet, headers: tuple, wide_column_index: int = None) -> None:
        for column_index, header in enumerate(headers, start=1):
            letter = get_column_letter(column_index)
            if column_index == wide_column_index:
                sheet.column_dimensions[letter].width = 60
                continue
            max_length = len(str(header))
            for row in sheet.iter_rows(min_col=column_index, max_col=column_index, min_row=2):
                for cell in row:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[letter].width = min(max_length + 2, 40)
