from typing import Any, Optional

from openpyxl import load_workbook

from model.actor import Actor
from model.block import Block
from model.character import Character

BLOCKS_SHEET_NAME = "Blocks"
REQUIRED_HEADERS = ("ID", "Part", "Scene", "Page", "Type", "Text")
CHARACTERS_SHEET_NAME = "Characters"
REQUIRED_CHARACTER_HEADERS = ("Mic", "Actor", "Character")


class WorkbookImporter:
    """Reads the blocks and the character roster back from a workbook written by WorkbookExporter.

    The workbook may have been edited by hand in the meantime, typically to fill
    in the `Audio Effect` and `Light` columns. Columns are therefore looked up by
    their header name, so their order does not matter and any additional columns
    are ignored.
    """

    def import_blocks(self, path: str) -> list[Block]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = self._sheet_rows(workbook, path, BLOCKS_SHEET_NAME)
            header = next(rows, None)
            if header is None:
                return []
            column_indices = self._column_indices(header, BLOCKS_SHEET_NAME, REQUIRED_HEADERS)
            return [self._to_block(row, column_indices) for row in rows if self._is_block_row(row, column_indices)]
        finally:
            workbook.close()

    def import_actors(self, path: str) -> list[Actor]:
        """Reads the actors with their microphone channel and characters from the Characters sheet.

        WorkbookExporter merges the Mic and Actor cells of an actor playing several
        characters, so these cells are empty in all but the first row of such an
        actor and are taken over from the row above.
        """
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = self._sheet_rows(workbook, path, CHARACTERS_SHEET_NAME)
            header = next(rows, None)
            if header is None:
                return []
            column_indices = self._column_indices(header, CHARACTERS_SHEET_NAME, REQUIRED_CHARACTER_HEADERS)
            actors: list[Actor] = []
            for row in rows:
                self._add_character_row(row, column_indices, actors)
            return actors
        finally:
            workbook.close()

    @staticmethod
    def _sheet_rows(workbook, path: str, sheet_name: str):
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"The workbook '{path}' does not contain a '{sheet_name}' sheet.")
        return workbook[sheet_name].iter_rows(values_only=True)

    @staticmethod
    def _column_indices(header: tuple, sheet_name: str, required_headers: tuple) -> dict[str, int]:
        column_indices = {str(name).strip(): index for index, name in enumerate(header) if name is not None}
        missing = [name for name in required_headers if name not in column_indices]
        if missing:
            raise ValueError(f"The '{sheet_name}' sheet is missing the columns: {', '.join(missing)}")
        return column_indices

    @staticmethod
    def _add_character_row(row: tuple, column_indices: dict[str, int], actors: list[Actor]) -> None:
        def cell(name: str) -> Any:
            return WorkbookImporter._cell(row, column_indices, name)

        character_name = WorkbookImporter._optional_text(cell("Character"))
        if character_name is None:
            return
        character = Character(name=character_name, actions=int(cell("Actions") or 0))

        mic = cell("Mic")
        actor_name = WorkbookImporter._optional_text(cell("Actor"))
        if mic is None and actor_name is None and actors:
            actors[-1].characters.append(character)
            return
        if mic is None or actor_name is None:
            raise ValueError(f"The character '{character_name}' in the '{CHARACTERS_SHEET_NAME}' sheet has no mic or actor.")
        actors.append(Actor(mic=int(mic), actor=actor_name, characters=[character]))

    @staticmethod
    def _is_block_row(row: tuple, column_indices: dict[str, int]) -> bool:
        return WorkbookImporter._cell(row, column_indices, "ID") is not None

    @staticmethod
    def _to_block(row: tuple, column_indices: dict[str, int]) -> Block:
        def cell(name: str) -> Any:
            return WorkbookImporter._cell(row, column_indices, name)

        return Block(
            id=int(cell("ID")),
            type=str(cell("Type") or ""),
            part=str(cell("Part") or ""),
            scene=int(cell("Scene") or 0),
            page=int(cell("Page") or 0),
            text=str(cell("Text") or ""),
            character=WorkbookImporter._optional_text(cell("Character")),
            direction=WorkbookImporter._optional_text(cell("Direction")),
            mute=WorkbookImporter._name_list(cell("Mute")),
            unmute=WorkbookImporter._name_list(cell("Unmute")),
            audio_effect=WorkbookImporter._optional_text(cell("Audio Effect")),
            light=WorkbookImporter._optional_text(cell("Light")),
        )

    @staticmethod
    def _cell(row: tuple, column_indices: dict[str, int], name: str) -> Any:
        index = column_indices.get(name)
        if index is None or index >= len(row):
            return None
        return row[index]

    @staticmethod
    def _optional_text(value: Any) -> Optional[str]:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _name_list(value: Any) -> list[str]:
        if value is None:
            return []
        return [name.strip() for name in str(value).split(",") if name.strip()]
